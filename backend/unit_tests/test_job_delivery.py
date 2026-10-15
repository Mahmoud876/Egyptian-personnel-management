"""Delivery/recovery and bounded export behavior using real SQLAlchemy queries and fakeredis."""
import json
from contextlib import contextmanager
from datetime import timedelta
from uuid import uuid4

import fakeredis
import pytest
from openpyxl import load_workbook
from redis.exceptions import ConnectionError
from rq.job import Job, JobStatus
from sqlalchemy import create_engine, insert
from sqlalchemy.dialects.mssql import DATETIME2
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import sessionmaker

from app.core.config import Settings
from app.core.security import utcnow
from app.recruits import jobs
from app.recruits.files import write_export
from app.recruits.models import ImportJob, ExportJob, ImportRecord, ImportRecordError
from app.recruits.validation import FIELDS


@compiles(DATETIME2, "sqlite")
def sqlite_datetime(type_, compiler, **kwargs):
    return "DATETIME"


@pytest.fixture
def delivery(monkeypatch, tmp_path):
    engine = create_engine("sqlite://")
    for model in (ImportJob, ExportJob, ImportRecord, ImportRecordError):
        model.__table__.create(engine)
    sessions = sessionmaker(engine, expire_on_commit=False)
    redis = fakeredis.FakeRedis()
    settings = Settings(db_user="test", db_password="test", storage_root=tmp_path,
                        dispatch_batch_size=1, export_batch_size=2)
    monkeypatch.setattr(jobs, "get_engine", lambda: engine)
    monkeypatch.setattr(jobs, "get_sessionmaker", lambda: sessions)
    monkeypatch.setattr(jobs, "get_redis", lambda: redis)
    monkeypatch.setattr(jobs, "get_settings", lambda: settings)
    yield sessions, redis, settings
    engine.dispose()


def new_import(sessions, **changes):
    with sessions() as db:
        item = ImportJob(original_filename="test.csv", storage_key="imports/test.csv", file_type="csv",
            file_sha256="0" * 64, batch_number=1, enlistment_year=2026, uploaded_by_user_id=1, **changes)
        db.add(item)
        db.commit()
        return item.id


def new_export(sessions, parent):
    with sessions() as db:
        item = ExportJob(import_job_id=parent, export_format="xlsx", requested_by_user_id=1)
        db.add(item)
        db.commit()
        return item.id


def test_targeted_delivery_is_idempotent_and_routes_exports(delivery):
    sessions, redis, settings = delivery
    parent = new_import(sessions)
    exported = new_export(sessions, parent)
    for _ in range(2):
        jobs.deliver_after_commit("import", str(parent), 0)
        jobs.deliver_after_commit("export", str(exported))
    assert jobs.queue().count == jobs.queue("export").count == 1
    with sessions() as db:
        for model, identity in ((ImportJob, parent), (ExportJob, exported)):
            item = db.get(model, identity)
            assert item.dispatched_at is not None
            assert item.retry_count == 0
            assert item.next_dispatch_check_at > utcnow() + timedelta(seconds=settings.queued_recheck_seconds - 5)


def test_stale_generation_and_terminal_jobs_are_not_delivered(delivery):
    sessions, _, _ = delivery
    parent = new_import(sessions, validation_run=2)
    terminal = new_import(sessions, status="CANCELLED")
    jobs.dispatch_one("import", str(parent), 1)
    jobs.dispatch_one("import", str(terminal), 0)
    assert jobs.queue().count == 0
    jobs.dispatch_one("import", str(parent), 2)
    assert jobs.queue().job_ids[0].endswith("-r2-1")


def test_lost_enqueue_ack_is_recovered_without_duplicate_delivery(delivery, monkeypatch):
    sessions, _, _ = delivery
    parent = new_import(sessions)
    real_queue = jobs.queue
    q = real_queue()
    enqueue = q.enqueue

    def lost_ack(*args, **kwargs):
        enqueue(*args, **kwargs)
        raise ConnectionError("synthetic lost acknowledgment")

    monkeypatch.setattr(q, "enqueue", lost_ack)
    monkeypatch.setattr(jobs, "queue", lambda kind="import": q)
    jobs.deliver_after_commit("import", str(parent), 0)
    assert q.count == 1
    with sessions() as db:
        item = db.get(ImportJob, parent)
        identity = item.rq_job_id
        assert item.dispatched_at is None and item.retry_count == 0
        item.next_dispatch_check_at = utcnow() - timedelta(seconds=1)
        db.commit()
    monkeypatch.setattr(jobs, "queue", real_queue)
    jobs.dispatch_pending()
    with sessions() as db:
        item = db.get(ImportJob, parent)
        assert item.rq_job_id == identity and item.dispatched_at is not None
    assert real_queue().count == 1


def test_due_pages_skip_future_terminal_and_busy_jobs_without_starvation(delivery, monkeypatch):
    sessions, _, _ = delivery
    busy = new_import(sessions)
    ready = new_import(sessions)
    future = new_import(sessions, next_dispatch_check_at=utcnow() + timedelta(hours=1))
    new_import(sessions, status="READY_FOR_REVIEW")
    actual_lock = jobs.worker_lock

    @contextmanager
    def lock(resource, wait_ms=0):
        if str(busy) in resource:
            yield False
        else:
            with actual_lock(resource, wait_ms) as acquired:
                yield acquired

    monkeypatch.setattr(jobs, "worker_lock", lock)
    jobs.dispatch_pending()
    with sessions() as db:
        assert db.get(ImportJob, busy).next_dispatch_check_at > utcnow()
        assert db.get(ImportJob, busy).rq_job_id is None
        assert db.get(ImportJob, ready).dispatched_at is not None
        assert db.get(ImportJob, future).rq_job_id is None
    assert jobs.queue().count == 1


def test_candidate_failure_does_not_block_next_page(delivery, monkeypatch):
    sessions, _, _ = delivery
    first = new_import(sessions)
    second = new_import(sessions)
    visited = []

    def attempt(kind, identity):
        visited.append(identity)
        if identity == str(first):
            raise RuntimeError("synthetic per-job failure")

    monkeypatch.setattr(jobs, "dispatch_one", attempt)
    jobs.dispatch_pending()
    assert visited == [str(first), str(second)]


def test_redis_outage_stops_scan_without_reserving_entire_backlog(delivery, monkeypatch):
    sessions, _, _ = delivery
    first = new_import(sessions)
    second = new_import(sessions)
    monkeypatch.setattr(jobs, "dispatch_one", lambda *args: (_ for _ in ()).throw(ConnectionError("outage")))
    jobs.dispatch_pending()
    with sessions() as db:
        assert db.get(ImportJob, first).next_dispatch_check_at > utcnow()
        assert db.get(ImportJob, second).next_dispatch_check_at < utcnow()


def test_legacy_export_keeps_its_original_queue(delivery):
    sessions, redis, _ = delivery
    parent = new_import(sessions)
    identity = new_export(sessions, parent)
    rq_id = f"export_job-{identity}-1"
    jobs.queue().enqueue(jobs.process_export, str(identity), job_id=rq_id)
    with sessions() as db:
        item = db.get(ExportJob, identity)
        item.rq_job_id = rq_id
        db.commit()
    jobs.dispatch_one("export", str(identity))
    assert jobs.queue().job_ids == [rq_id]
    assert jobs.queue("export").count == 0


def test_prestart_failures_remain_bounded(delivery):
    sessions, redis, settings = delivery
    identity = new_import(sessions)
    for attempt in range(1, settings.job_max_attempts + 1):
        jobs.dispatch_one("import", str(identity))
        with sessions() as db:
            item = db.get(ImportJob, identity)
            assert item.rq_job_id.endswith(f"-{attempt}")
            rq_job = Job.fetch(item.rq_job_id, connection=redis)
            jobs.queue().remove(rq_job.id)
            rq_job.set_status(JobStatus.FAILED)
    jobs.dispatch_one("import", str(identity))
    with sessions() as db:
        item = db.get(ImportJob, identity)
        assert item.status == "FAILED" and item.retry_count == settings.job_max_attempts


def seed_export(sessions, parent, count):
    with sessions() as db:
        for number in range(1, count + 1):
            original = {field: f"{field}-{number}" for field in FIELDS}
            original["name"] = "أحمد محمد"
            original["rejection_reasons"] = "original note"
            if number == count:
                original["source_rejection_reasons"] = "reserved late"
                original["original_uploaded_data"] = "reserved original column"
                original["late_column"] = "late value"
            corrected = {**original, "name": "=literal"} if number == 1 else None
            identity = uuid4()
            db.execute(insert(ImportRecord), dict(id=identity, import_job_id=parent, row_number=number,
                raw_data_json=json.dumps(original, ensure_ascii=False),
                corrected_data_json=json.dumps(corrected, ensure_ascii=False) if corrected else None,
                validation_status="INVALID"))
            db.execute(insert(ImportRecordError), dict(id=uuid4(), import_row_id=identity,
                error_code="ID_LENGTH", field_name="national_id", error_message="سبب الرفض"))
        db.commit()


def test_export_streams_across_batches_and_preserves_late_headers(delivery, tmp_path):
    sessions, _, _ = delivery
    parent = new_import(sessions)
    seed_export(sessions, parent, 5)
    output = tmp_path / "export.xlsx"
    with sessions() as db:
        headers, reasons, original = jobs.export_layout(db, parent)
        assert reasons == "source_source_rejection_reasons"
        assert original == "source_original_uploaded_data"
        assert "late_column" in headers
        rows = jobs.export_rows(db, parent, reasons, original)
        assert iter(rows) is rows
        assert write_export(output, headers, rows) == 5
    book = load_workbook(output)
    sheet = book.active
    assert sheet.max_row == 6
    first = dict(zip(headers, next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))))
    assert first[reasons] == "original note"
    assert json.loads(first[original])["name"] == "أحمد محمد"
    assert "سبب الرفض" in first["rejection_reasons"]
    assert sheet.cell(2, headers.index("name") + 1).data_type == "s"
    assert sheet.cell(6, headers.index("late_column") + 1).value == "late value"
    book.close()


def test_empty_export_still_has_headers(delivery, tmp_path):
    sessions, _, _ = delivery
    parent = new_import(sessions)
    with sessions() as db:
        headers, reasons, original = jobs.export_layout(db, parent)
        path = tmp_path / "empty.xlsx"
        assert write_export(path, headers, jobs.export_rows(db, parent, reasons, original)) == 0
    book = load_workbook(path)
    assert list(book.active.values) == [tuple(FIELDS) + ("rejection_reasons",)]
    book.close()


def test_large_export_memory_stays_bounded_in_fresh_processes(tmp_path):
    import os
    import subprocess
    import sys
    from pathlib import Path
    from app.core.config import REPO_ROOT
    measurements = []
    for count in (10_000, 100_000):
        root = tmp_path / str(count)
        root.mkdir()
        result = subprocess.run([sys.executable, str(Path(__file__).with_name("export_memory_probe.py")),
            str(count), str(root)], cwd=REPO_ROOT / "backend", env=dict(os.environ, PYTHONPATH=str(REPO_ROOT / "backend")),
            capture_output=True, text=True, timeout=180, check=True)
        measurement = json.loads(result.stdout)
        assert measurement["rows"] == count
        measurements.append(measurement)
    print("Export memory:", json.dumps(measurements))
    assert measurements[1]["python_peak_bytes"] < max(8 * 1024 * 1024, measurements[0]["python_peak_bytes"] * 3)
