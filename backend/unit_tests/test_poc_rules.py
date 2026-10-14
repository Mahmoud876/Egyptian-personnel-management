"""Pure rules and file/session contracts; no SQL Server or real personal data."""
import csv
import json
from datetime import date
from io import BytesIO, StringIO
from types import SimpleNamespace

import fakeredis
import pytest
from openpyxl import load_workbook, Workbook

from app.core import security
from app.core.roles import permissions_for_role
from app.recruits.validation import validate_row, add_months, FIELDS
from app.recruits.files import read_records, write_export, path_for


def valid_row(**changes):
    return {"id": "1", "name": "أحمد محمد حسن علي", "national_id": "30001010100110",
            "birth_date": "2000-01-01", "gov_id": "01", "gov_name": "القاهرة",
            "ps_id": "0101", "ps_name": "قسم تجريبي", "type_id": "1", "type_name": "عالى", **changes}


@pytest.mark.parametrize("changes,code", [
    ({"national_id": "1"}, "ID_LENGTH"),
    ({"national_id": "X" * 14}, "ID_DIGITS"),
    ({"national_id": "30002300100110"}, "ID_BIRTH_DATE"),
    ({"national_id": "40001010100110"}, "ID_BIRTH_DATE"),
    ({"birth_date": "2001-01-01"}, "BIRTH_MISMATCH"),
    ({"national_id": "30001010100120"}, "GENDER"),
    ({"gov_id": "99"}, "GOV_UNKNOWN"),
    ({"gov_id": "21"}, "GOV_MISMATCH"),
    ({"type_id": "5"}, "QUALIFICATION"),
    ({"name": ""}, "FIELD_REQUIRED"),
    ({"birth_date": "not a date"}, "BIRTH_DATE"),
])
def test_all_validation_cases(changes, code):
    _, errors = validate_row(valid_row(**changes), set(), False, {1, 2, 3, 4})
    assert code in {e["error_code"] for e in errors}


def test_valid_official_giza_code_and_all_reasons():
    _, errors = validate_row(valid_row(national_id="30001012100110", gov_id="21", gov_name="الجيزة"), set(), False, {1})
    assert errors == []
    _, errors = validate_row(valid_row(name="", birth_date="bad"), {"30001010100110"}, True, {1})
    assert {"FIELD_REQUIRED", "BIRTH_DATE", "DUPLICATE_FILE", "DUPLICATE_DATABASE"} <= {e["error_code"] for e in errors}


@pytest.mark.parametrize("months,expected", [(14, date(2027, 3, 1)), (20, date(2027, 9, 1)), (26, date(2028, 3, 1)), (38, date(2029, 3, 1))])
def test_service_dates(months, expected):
    assert add_months(date(2026, 1, 1), months) == expected


def test_month_end_and_leap_year():
    assert add_months(date(2024, 2, 29), 12) == date(2025, 2, 28)
    assert add_months(date(2026, 12, 31), 2) == date(2027, 2, 28)


def test_csv_and_xlsx_preserve_raw_values(tmp_path):
    path = tmp_path / "sample.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerow(valid_row())
    assert read_records(path) == [valid_row()]
    book = Workbook()
    book.active.append(list(FIELDS))
    book.active.append([valid_row()[h] for h in FIELDS])
    path = tmp_path / "sample.xlsx"
    book.save(path)
    assert read_records(path) == [valid_row()]


def test_export_formula_text_is_not_executable(tmp_path):
    path = tmp_path / "rejected.xlsx"
    write_export(path, [{**valid_row(name="=HYPERLINK(\"https://invalid\")"), "rejection_reasons": "سبب الرفض"}])
    book = load_workbook(path)
    assert book.active["B2"].data_type == "s"
    assert book.active["B2"].value.startswith("=HYPERLINK")
    assert book.active.cell(2, 11).value == "سبب الرفض"


def test_missing_duplicate_headers_and_row_limit(tmp_path, monkeypatch):
    path = tmp_path / "bad.csv"
    path.write_text("id,id\n1,2", encoding="utf-8")
    with pytest.raises(ValueError, match="INVALID_COLUMNS"):
        read_records(path)


def test_storage_traversal_is_rejected():
    with pytest.raises(ValueError):
        path_for("../../outside.txt")


def test_roles_are_closed_and_operator_can_complete():
    operator = permissions_for_role("data_entry_operator")
    assert {"profile:get", "profile:insert", "barcode:get"} <= operator
    assert not {"file:insert", "profile:update", "audit:get", "user:update"} & operator
    assert permissions_for_role("invented_role") == set()


def test_redis_sessions_are_hashed_expire_and_logout(monkeypatch):
    redis = fakeredis.FakeRedis()
    monkeypatch.setattr(security, "get_redis", lambda: redis)
    settings = SimpleNamespace(session_max_hours=12, session_idle_minutes=30)
    monkeypatch.setattr(security, "get_settings", lambda: settings)
    db = SimpleNamespace(get=lambda *args: SimpleNamespace(session_version=2))
    token = security.create_session(db, 7, None)
    session = security.resolve_session(db, token)
    assert session.user_id == 7 and session.version == 2
    assert token.encode() not in b" ".join(redis.keys())
    redis_key = redis.keys()[0]
    assert 0 < redis.ttl(redis_key) <= 1800
    data = json.loads(redis.get(redis_key))
    data["expires"] = 0
    redis.set(redis_key, json.dumps(data))
    assert security.resolve_session(db, token) is None
    token = security.create_session(db, 7, None)
    security.delete_session(db, token)
    assert security.resolve_session(db, token) is None
    assert security.resolve_session(db, "é" * 64) is None
