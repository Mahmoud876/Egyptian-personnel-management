"""PoC contract tests against SQL Server and Redis, with synthetic data only."""
import csv
import json
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from io import StringIO, BytesIO
from itertools import count
from uuid import UUID, uuid4

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select, text

from app.core.config import get_settings
from app.core.db import get_sessionmaker
from app.core.models import AppUser, AuditLog
from app.core.security import hash_password
from app.core.redis_client import get_redis
from app.main import create_app
from app.recruits.jobs import process_import, process_export, dispatch_pending, queue
from app.recruits.models import ImportJob, Recruit, ImportRecord, ImportRecordError, ExportJob
from app.recruits.validation import FIELDS

PASSWORD = "Synthetic-test-42!"
SYNTHETIC_SERIALS = count(1, 2)


@pytest.fixture
def actors():
    clients = []
    for role in ["system_administrator", "data_entry_operator"]:
        username = "poc_" + uuid4().hex[:16]
        with get_sessionmaker()() as db:
            user = AppUser(username=username, full_name="مستخدم تجريبي", role=role,
                           password_hash=hash_password(PASSWORD), is_active=True, must_change_password=False)
            db.add(user)
            db.commit()
        client = TestClient(create_app())
        response = client.post("/api/v1/auth/login", json={"username": username, "password": PASSWORD})
        assert response.status_code == 200, response.text
        assert response.json()["role"] == role
        clients.append(client)
    yield clients
    for client in clients:
        client.close()


def row(nid=None, **changes):
    # Unique within this test session; random IDs could collide with approved recruits.
    serial = next(SYNTHETIC_SERIALS)
    assert serial <= 9999, "Synthetic national-ID serial range exhausted"
    return {"id": "1", "name": "أحمد محمد حسن علي", "national_id": nid or f"300010101{serial:04d}0",
            "birth_date": "2000-01-01", "gov_id": "01", "gov_name": "القاهرة",
            "ps_id": "0101", "ps_name": "قسم تجريبي", "type_id": "1", "type_name": "عالى", **changes}


def upload(client, rows):
    stream = StringIO()
    writer = csv.DictWriter(stream, fieldnames=list(FIELDS) + sorted({key for item in rows for key in item} - set(FIELDS)))
    writer.writeheader()
    writer.writerows(rows)
    response = client.post("/api/v1/imports", data={"batch_number": "1", "enlistment_year": "2026"},
                           files={"file": ("synthetic.csv", stream.getvalue().encode("utf-8"), "text/csv")})
    assert response.status_code == 202, response.text
    return response.json()["id"]


def approved_recruit(admin):
    job = upload(admin, [row()])
    process_import(job)
    response = admin.post(f"/api/v1/imports/{job}/approve")
    assert response.status_code == 200, response.text
    with get_sessionmaker()() as db:
        return str(db.scalar(select(Recruit.id).where(Recruit.import_job_id == UUID(job))))


def completion(**changes):
    return {"police_number": uuid4().hex[:16], "enlistment_start_date": "2026-01-01",
            "enlistment_end_date": "2027-03-01", **changes}


def test_import_status_counts_and_official_reference(actors):
    admin, operator = actors
    first = upload(admin, [row()])
    second = upload(admin, [row()])
    process_import(first)
    for statuses in [["READY_FOR_REVIEW"], ["QUEUED", "PARSING", "VALIDATING"]]:
        with get_sessionmaker()() as db:
            expected = len(list(db.scalars(select(ImportJob.id).where(ImportJob.status.in_(statuses)))))
        response = admin.get("/api/v1/imports", params=[("page_size", 1)] + [("status", s) for s in statuses])
        assert response.status_code == 200, response.text
        assert response.json()["total"] == expected
        assert all(item["status"] in statuses for item in response.json()["items"])
    assert admin.get("/api/v1/imports?status=UNKNOWN").status_code == 422
    assert operator.get("/api/v1/imports?status=QUEUED").status_code == 403
    references = operator.get("/api/v1/lookups/official-governorates")
    assert references.status_code == 200
    names = {item["code"]: item["name_ar"] for item in references.json()}
    assert len(names) == 28
    assert names["21"] == "الجيزة" and names["11"] == "دمياط"
    # Leave no pending import for a later worker test to pick up.
    assert admin.post(f"/api/v1/imports/{first}/cancel").status_code == 200
    process_import(second)
    assert admin.post(f"/api/v1/imports/{second}/cancel").status_code == 200


def test_import_review_approve_complete_lock_print_search(actors):
    admin, operator = actors
    good = row()
    job = upload(admin, [good, row(nid="bad")])
    process_import(job)
    report = admin.get(f"/api/v1/imports/{job}/report").json()
    assert report["summary"]["status"] == "READY_FOR_REVIEW"
    assert report["summary"]["valid_rows"] == 1
    assert report["invalid"]["total"] == 1
    assert len(report["invalid"]["items"][0]["errors"]) >= 2
    assert admin.post(f"/api/v1/imports/{job}/approve").json()["inserted_count"] == 1
    assert admin.post(f"/api/v1/imports/{job}/approve").status_code == 409
    recruit = operator.get("/api/v1/recruits/by-national-id/" + good["national_id"]).json()
    rid = recruit["id"]
    assert recruit["suggested_end_date"] == "2027-03-01"
    assert operator.get(f"/api/v1/recruits/{rid}/barcode").status_code == 409
    assert operator.post(f"/api/v1/recruits/{rid}/complete", json=completion(enlistment_end_date="2027-02-28")).status_code == 422
    assert operator.post(f"/api/v1/recruits/{rid}/complete", json=completion(enlistment_start_date="9999-12-31", enlistment_end_date="9999-12-31")).status_code == 422
    body = completion()
    done = operator.post(f"/api/v1/recruits/{rid}/complete", json=body)
    assert done.status_code == 200, done.text
    assert done.json()["is_locked"] and len(done.json()["barcode"]) == 12
    assert operator.post(f"/api/v1/recruits/{rid}/complete", json=body).status_code == 409
    assert operator.patch(f"/api/v1/recruits/{rid}", json={**body, "revision": 2}).status_code == 403
    changed = admin.patch(f"/api/v1/recruits/{rid}", json={**body, "revision": 2, "name": "محمد علي حسن أحمد"})
    assert changed.status_code == 200, changed.text
    assert changed.json()["barcode"] == done.json()["barcode"]
    label = operator.get(f"/api/v1/recruits/{rid}/barcode")
    assert label.status_code == 200 and "<svg" in label.text and body["police_number"] in label.text
    assert operator.get("/api/v1/recruits", params={"barcode": done.json()["barcode"]}).json()["total"] == 1
    assert operator.get("/api/v1/recruits", params={"police_number": body["police_number"]}).status_code == 422
    assert admin.get("/api/v1/audit", params={"q": good["national_id"]}).json()["total"] >= 1


def test_operator_denials_and_cancel(actors):
    admin, operator = actors
    for path in ["/imports", "/audit", "/users"]:
        assert operator.get("/api/v1" + path).status_code == 403
    job = upload(admin, [row()])
    assert admin.post(f"/api/v1/imports/{job}/approve").status_code == 409
    process_import(job)
    assert admin.post(f"/api/v1/imports/{job}/cancel").status_code == 200
    assert admin.post(f"/api/v1/imports/{job}/approve").status_code == 409
    with get_sessionmaker()() as db:
        assert not db.scalar(select(Recruit.id).where(Recruit.import_job_id == UUID(job)))


def test_duplicate_file_and_database_recheck_on_approval(actors):
    admin, _ = actors
    same = row()
    duplicate_job = upload(admin, [same, same])
    process_import(duplicate_job)
    assert admin.get(f"/api/v1/imports/{duplicate_job}").json()["invalid_rows"] == 2
    first, second = upload(admin, [same]), upload(admin, [same])
    process_import(first)
    process_import(second)
    assert admin.post(f"/api/v1/imports/{first}/approve").status_code == 200
    assert admin.post(f"/api/v1/imports/{second}/approve").status_code == 409
    assert admin.get(f"/api/v1/imports/{second}").json()["status"] == "READY_FOR_REVIEW"
    version = admin.get(f"/api/v1/imports/{second}").json()["progress_version"]
    assert admin.post(f"/api/v1/imports/{second}/revalidate", json={"version": version}).status_code == 202
    process_import(second)
    report = admin.get(f"/api/v1/imports/{second}/report").json()
    assert report["summary"]["valid_rows"] == 0
    assert any(error["error_code"] == "DUPLICATE_DATABASE" for error in report["invalid"]["items"][0]["errors"])


def test_correction_preserves_source_rechecks_all_rows_and_locks_review(actors):
    admin, operator = actors
    original = row()
    job = upload(admin, [original, original])
    process_import(job)
    report = admin.get(f"/api/v1/imports/{job}/report").json()
    target = report["invalid"]["items"][0]
    version = report["summary"]["progress_version"]
    body = {"version": version, "changes": {"national_id": row()["national_id"]}}
    path = f"/api/v1/imports/{job}/rows/{target['id']}"
    assert operator.patch(path, json=body).status_code == 403
    assert admin.patch(f"/api/v1/imports/{job}/rows/{uuid4()}", json=body).status_code == 404
    assert admin.patch(path, json={**body, "version": version - 1}).status_code == 409
    assert admin.patch(path, json={**body, "changes": {"validation_status": "VALID"}}).status_code == 422
    assert admin.patch(path, json=body).status_code == 202
    assert admin.post(f"/api/v1/imports/{job}/approve").status_code == 409
    assert admin.patch(path, json=body).status_code == 409
    assert admin.post(f"/api/v1/imports/{job}/exports", json={"format": "xlsx"}).status_code == 409
    process_import(job)
    report = admin.get(f"/api/v1/imports/{job}/report").json()
    assert report["summary"]["valid_rows"] == 2 and report["summary"]["invalid_rows"] == 0
    corrected = next(item for item in report["valid"]["items"] if item["id"] == target["id"])
    assert corrected["raw_data"] == original
    assert corrected["corrected_data"]["national_id"] == body["changes"]["national_id"]
    assert corrected["errors"] == []
    assert admin.patch(path, json=body).status_code == 409  # old review remains stale after revalidation
    with get_sessionmaker()() as db:
        log = db.scalar(select(AuditLog).where(AuditLog.entity_id == target["id"], AuditLog.action == "import_row_corrected"))
        assert log is not None
    assert admin.post(f"/api/v1/imports/{job}/approve").json()["inserted_count"] == 2
    assert admin.patch(path, json={**body, "version": report["summary"]["progress_version"]}).status_code == 409


def test_correction_export_keeps_current_and_original_values(actors):
    from openpyxl import load_workbook
    admin, _ = actors
    job = upload(admin, [row(nid="bad", name="")])
    process_import(job)
    report = admin.get(f"/api/v1/imports/{job}/report").json()
    target = report["invalid"]["items"][0]
    path = f"/api/v1/imports/{job}/rows/{target['id']}"
    body = {"version": report["summary"]["progress_version"], "changes": {"name": "اسم مصحح من المصدر"}}
    export = admin.post(f"/api/v1/imports/{job}/exports", json={"format": "xlsx"}).json()["id"]
    assert admin.patch(path, json=body).json()["code"] == "export_in_progress"
    process_export(export)
    assert admin.patch(path, json=body).status_code == 202
    process_import(job)
    report = admin.get(f"/api/v1/imports/{job}/report").json()
    assert report["summary"]["invalid_rows"] == 1
    assert not any(error["field_name"] == "name" for error in report["invalid"]["items"][0]["errors"])
    export = admin.post(f"/api/v1/imports/{job}/exports", json={"format": "xlsx"}).json()["id"]
    process_export(export)
    book = load_workbook(BytesIO(admin.get(f"/api/v1/exports/{export}/download").content))
    values = list(book.active.values)
    data = dict(zip(values[0], values[1]))
    assert data["name"] == body["changes"]["name"]
    assert json.loads(data["original_uploaded_data"])["name"] == ""
    assert data["rejection_reasons"]


def test_rejection_reasons_distinguish_missing_and_length_in_existing_reports_and_exports(actors):
    from openpyxl import load_workbook
    admin, _ = actors
    job = upload(admin, [row(ps_name="", gov_name="", name="س" * 201)])
    process_import(job)
    report = admin.get(f"/api/v1/imports/{job}/report").json()
    target = report["invalid"]["items"][0]
    errors = {e["field_name"]: e for e in target["errors"]}
    assert errors["name"]["error_code"] == "FIELD_TOO_LONG"
    assert "200" in errors["name"]["error_message"] and "201" in errors["name"]["error_message"]
    assert errors["ps_name"]["error_code"] == "FIELD_REQUIRED"
    # Simulate a report persisted before this fix: it must improve without revalidation.
    with get_sessionmaker()() as db:
        for error in db.scalars(select(ImportRecordError).where(ImportRecordError.import_row_id == UUID(target["id"]))):
            error.error_code = "FIELD_REQUIRED"
            error.error_message = "البيان مطلوب أو يتجاوز الطول المسموح"
        db.commit()
    current = admin.get(f"/api/v1/imports/{job}/rows?status=INVALID").json()["items"][0]
    assert {e["field_name"]: e for e in current["errors"]} == errors
    export = admin.post(f"/api/v1/imports/{job}/exports", json={"format": "xlsx"}).json()["id"]
    process_export(export)
    sheet = load_workbook(BytesIO(admin.get(f"/api/v1/exports/{export}/download").content)).active
    values = list(sheet.values)
    result = dict(zip(values[0], values[1]))
    assert result["name"] == "س" * 201
    for field in ["ps_name", "gov_name", "name"]:
        assert f"({field}):" in result["rejection_reasons"]
    assert "200" in result["rejection_reasons"] and "201" in result["rejection_reasons"]


def test_validation_completion_uses_the_requester_of_each_run(actors, make_user):
    admin, _ = actors
    uploader = admin.get("/api/v1/auth/me").json()["id"]
    second_user = make_user(role="system_administrator")
    with TestClient(create_app()) as second:
        assert second.post("/api/v1/auth/login", json={"username": second_user.username,
            "password": second_user.password}).status_code == 200
        job = upload(admin, [row(nid="bad")])
        process_import(job)
        report = admin.get(f"/api/v1/imports/{job}/report").json()
        target = report["invalid"]["items"][0]
        assert second.patch(f"/api/v1/imports/{job}/rows/{target['id']}", json={
            "version": report["summary"]["progress_version"], "changes": {"name": "تصحيح المدير الثاني"}}).status_code == 202
        # Recovery/retries must keep the initiating administrator, not revert to the uploader.
        with get_sessionmaker()() as db:
            item = db.get(ImportJob, UUID(job))
            item.retry_count = 1
            db.commit()
        process_import(job)
        version = admin.get(f"/api/v1/imports/{job}").json()["progress_version"]
        assert admin.post(f"/api/v1/imports/{job}/revalidate", json={"version": version}).status_code == 202
        process_import(job)
        with get_sessionmaker()() as db:
            entries = list(db.scalars(select(AuditLog).where(AuditLog.entity_id == job,
                AuditLog.action == "validation_complete").order_by(AuditLog.id)))
            assert [entry.user_id for entry in entries] == [uploader, second_user.id, uploader]
            assert [json.loads(entry.new_values)["validation_run"] for entry in entries] == [0, 1, 2]


def test_partial_identifier_search_returns_all_matches_with_paging_and_literal_wildcards(actors):
    admin, operator = actors
    job = upload(admin, [row(nid="30001010111310", id="1"), row(nid="30001010121310", id="2")])
    process_import(job)
    assert admin.post(f"/api/v1/imports/{job}/approve").status_code == 200
    with get_sessionmaker()() as db:
        ids = list(db.scalars(select(Recruit.id).where(Recruit.import_job_id == UUID(job)).order_by(Recruit.national_id)))
    profiles = []
    for index, rid in enumerate(ids):
        response = operator.post(f"/api/v1/recruits/{rid}/complete", json=completion(police_number=f"{index}PART%_5678"))
        assert response.status_code == 200, response.text
        profiles.append(response.json())
    for params in [{"national_id": "1310"}, {"search": "1310"}, {"search": "PART%_"},
                   {"police_number": "5678", "enlistment_year": 2026}]:
        result = operator.get('/api/v1/recruits', params={**params, 'page_size': 1}).json()
        assert result['total'] == 2
        other = operator.get('/api/v1/recruits', params={**params, 'page_size': 1, 'page': 2}).json()
        assert {result['items'][0]['id'], other['items'][0]['id']} == {str(rid) for rid in ids}
    assert operator.get('/api/v1/recruits', params={'police_number': '5678', 'enlistment_year': 2025}).json()['total'] == 0
    assert operator.get('/api/v1/recruits', params={'police_number': '5678'}).status_code == 422
    for profile in profiles:
        for key in ['barcode', 'search']:
            response = operator.get('/api/v1/recruits', params={key: profile['barcode'][-4:]})
            assert response.status_code == 200
            assert profile['id'] in {item['id'] for item in response.json()['items']}
    assert operator.get('/api/v1/recruits', params={'national_id': '1%'}).status_code == 422
    assert operator.get('/api/v1/recruits', params={'barcode': '1' * 13}).status_code == 422
    # Exact lookup is still exact; partial searching cannot weaken identity validation.
    assert operator.get('/api/v1/recruits/by-national-id/1310').status_code == 422
    assert operator.get('/api/v1/recruits', params={'search': '%[_]'}).json()['total'] == 0


def test_audit_search_finds_arabic_actions_and_related_recruit_events(actors):
    admin, operator = actors
    rid = approved_recruit(admin)
    unrelated = approved_recruit(admin)
    assert operator.post(f"/api/v1/recruits/{rid}/complete", json=completion()).status_code == 200
    assert operator.get(f"/api/v1/recruits/{rid}/barcode").status_code == 200
    profile = operator.get(f"/api/v1/recruits/{rid}").json()
    assert admin.post(f"/api/v1/recruits/{rid}/documents", files={
        "file": ("synthetic.pdf", b"%PDF-synthetic", "application/pdf")}).status_code == 201
    result = admin.get("/api/v1/audit", params={"q": profile["national_id"], "page_size": 100})
    assert result.status_code == 200, result.text
    entries = result.json()["items"]
    assert {"label", "complete", "document_upload", "approve", "validation_complete"} <= {e["action"] for e in entries}
    assert all(e["entity_id"] != unrelated for e in entries)
    for query, action in [("طباعة ملصق", "label"), ("إعتماد", "approve"), ("إرفاق مستند", "document_upload")]:
        result = admin.get("/api/v1/audit", params={"q": query})
        assert result.status_code == 200, result.text
        assert action in {e["action"] for e in result.json()["items"]}
    # Exact filters still narrow global search and wildcard characters stay literal.
    result = admin.get("/api/v1/audit", params={"q": profile["national_id"], "action": "label"}).json()
    assert result["total"] == 1 and result["items"][0]["entity_id"] == rid
    assert admin.get("/api/v1/audit", params={"q": "%[_]"}).json()["total"] == 0


@pytest.mark.parametrize("status", ["QUEUED", "PARSING", "VALIDATING", "FAILED"])
def test_revalidation_rows_remain_readable_but_review_actions_are_locked(actors, status):
    admin, operator = actors
    original = row(nid="bad")
    job = upload(admin, [original, row()])
    base = f"/api/v1/imports/{job}"
    # An initial upload has no completed staging report to inspect yet.
    assert admin.get(base + "/rows").status_code == 409
    process_import(job)
    report = admin.get(base + "/report").json()
    target = report["invalid"]["items"][0]
    correction = {"version": report["summary"]["progress_version"], "changes": {"name": "اسم مصحح"}}
    assert admin.patch(base + f"/rows/{target['id']}", json=correction).status_code == 202
    with get_sessionmaker()() as db:
        item = db.get(ImportJob, UUID(job))
        item.status = status
        db.commit()
    result = admin.get(base + "/rows?status=INVALID&page_size=1")
    assert result.status_code == 200, result.text
    record = result.json()["items"][0]
    assert record["raw_data"] == original
    assert record["corrected_data"]["name"] == "اسم مصحح"
    assert record["errors"]  # explicitly provisional until validation completes
    assert admin.get(base + "/rows?status=VALID").json()["total"] == 1
    assert admin.get(base + "/report").status_code == 200
    assert operator.get(base + "/rows").status_code == 403
    current = admin.get(base).json()["progress_version"]
    assert admin.patch(base + f"/rows/{target['id']}", json={**correction, "version": current}).status_code == 409
    assert admin.post(base + "/approve").status_code == 409
    assert admin.post(base + "/exports", json={"format": "xlsx"}).status_code == 409
    if status == "FAILED":
        assert admin.post(base + "/revalidate", json={"version": current}).status_code == 202
    process_import(job)
    assert admin.get(base).json()["status"] == "READY_FOR_REVIEW"


def test_export_and_terminal_sse(actors):
    admin, _ = actors
    job = upload(admin, [row(nid="bad", rejection_reasons="original note", source_rejection_reasons="keep this too")])
    process_import(job)
    stream = admin.get(f"/api/v1/imports/{job}/events")
    assert stream.status_code == 200 and "READY_FOR_REVIEW" in stream.text
    response = admin.post(f"/api/v1/imports/{job}/exports", json={"format": "xlsx"})
    assert response.status_code == 202
    export = response.json()["id"]
    assert admin.get(f"/api/v1/exports/{export}/download").status_code == 409
    process_export(export)
    result = admin.get(f"/api/v1/exports/{export}/download")
    assert result.status_code == 200 and result.content[:2] == b"PK"
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(result.content), read_only=True)
    values = list(book.active.values)
    exported = dict(zip(values[0], values[1]))
    assert exported["source_source_rejection_reasons"] == "original note"
    assert exported["source_rejection_reasons"] == "keep this too"
    assert "الرقم القومي" in exported["rejection_reasons"]
    book.close()


def test_restart_does_not_duplicate_staging(actors):
    admin, _ = actors
    job = upload(admin, [row() for _ in range(260)])
    process_import(job)
    with get_sessionmaker()() as db:
        item = db.get(ImportJob, UUID(job))
        item.status = "VALIDATING"  # crash after checkpoint commit, before final state commit
        db.commit()
    process_import(job)
    with get_sessionmaker()() as db:
        records = db.scalars(select(ImportRecord).where(ImportRecord.import_job_id == UUID(job))).all()
        assert len(records) == 260
    assert admin.get(f"/api/v1/imports/{job}").json()["processed_rows"] == 260


def test_concurrent_profile_completion_has_one_winner(actors):
    admin, operator = actors
    recruit = approved_recruit(admin)
    def save(_):
        with TestClient(create_app()) as client:
            client.cookies.update(operator.cookies)
            return client.post(f"/api/v1/recruits/{recruit}/complete", json=completion()).status_code
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(save, range(2))) == [200, 409]


def test_role_changes_and_deactivation_invalidate_sessions(actors):
    admin, operator = actors
    user = operator.get("/api/v1/auth/me").json()
    response = admin.put(f'/api/v1/users/{user["id"]}', json={"full_name": user["full_name"],
        "work_entity_id": None, "is_active": True, "role": "system_administrator"})
    assert response.status_code == 200, response.text
    assert operator.get("/api/v1/auth/me").status_code == 401
    response = admin.post("/api/v1/users", json={"username": "new_" + uuid4().hex[:12], "full_name": "اختبار",
        "password": PASSWORD, "role": "data_entry_operator", "permissions": ["file:insert"]})
    assert response.status_code == 422


def test_dispatcher_delivers_sql_job_to_real_rq(actors):
    from rq import SimpleWorker
    from rq.serializers import JSONSerializer
    from rq.timeouts import TimerDeathPenalty
    class TestWorker(SimpleWorker):
        death_penalty_class = TimerDeathPenalty
    admin, _ = actors
    job = upload(admin, [row()])
    dispatch_pending()
    TestWorker([queue()], connection=get_redis(), serializer=JSONSerializer).work(burst=True)
    assert admin.get(f"/api/v1/imports/{job}").json()["status"] == "READY_FOR_REVIEW"


def test_worker_crash_after_checkpoint_resumes_without_duplicates(actors, monkeypatch):
    from app.recruits import jobs
    admin, _ = actors
    job = upload(admin, [row() for _ in range(270)])
    original = jobs.publish
    def crash(item):
        if item.processed_rows == 250 and item.status == "VALIDATING":
            raise RuntimeError("simulated process failure after checkpoint")
        original(item)
    monkeypatch.setattr(jobs, "publish", crash)
    with pytest.raises(RuntimeError):
        process_import(job)
    assert admin.get(f"/api/v1/imports/{job}").json()["checkpoint_row"] == 250
    monkeypatch.setattr(jobs, "publish", original)
    process_import(job)
    state = admin.get(f"/api/v1/imports/{job}").json()
    assert state["processed_rows"] == 270 and state["retry_count"] == 2
    with get_sessionmaker()() as db:
        assert len(db.scalars(select(ImportRecord).where(ImportRecord.import_job_id == UUID(job))).all()) == 270


def test_parsing_publishes_persisted_progress_before_eof_and_checks_cross_batch_duplicates(actors, monkeypatch):
    from app.recruits import jobs
    admin, _ = actors
    # Distinct synthetic IDs, with a duplicate crossing the 250-row evaluation boundary.
    rows = [row(nid=f"301020202{2 * i + 1:04d}0", birth_date="2001-02-02", gov_id="02", gov_name="الإسكندرية") for i in range(511)]
    rows.append(rows[0])
    job = upload(admin, rows)
    consumed = 0
    real_iter = jobs.iter_records
    def tracked(path):
        nonlocal consumed
        for record in real_iter(path):
            consumed += 1
            yield record
    events = []
    def capture(item):
        events.append(jobs.snapshot(item))
        if item.status == "PARSING" and item.parsed_rows == 250:
            assert consumed == 250  # discovery is not waiting for full materialization
            persisted = admin.get(f"/api/v1/imports/{job}").json()
            assert persisted["parsed_rows"] == 250 and persisted["total_rows"] == 250
            assert persisted["processed_rows"] == 0
    monkeypatch.setattr(jobs, "iter_records", tracked)
    monkeypatch.setattr(jobs, "publish", capture)
    process_import(job)
    assert [e["parsed_rows"] for e in events if e["status"] == "PARSING"] == [0, 250, 500, 512]
    assert [e["processed_rows"] for e in events if e["status"] == "VALIDATING"] == [0, 250, 500, 512]
    assert all(e["processed_rows"] == e["valid_rows"] + e["invalid_rows"] for e in events)
    assert events[-1]["invalid_rows"] == 2
    assert events[-1]["valid_rows"] == 510


def test_parsing_restart_resets_read_counter_without_creating_duplicate_staging(actors, monkeypatch):
    from app.recruits import jobs
    admin, _ = actors
    job = upload(admin, [row() for _ in range(260)])
    original = jobs.iter_records
    def broken(path):
        for number, record in enumerate(original(path)):
            if number == 250:
                raise RuntimeError("synthetic reader interruption")
            yield record
    monkeypatch.setattr(jobs, "iter_records", broken)
    with pytest.raises(RuntimeError):
        process_import(job)
    state = admin.get(f"/api/v1/imports/{job}").json()
    assert state["parsed_rows"] == 250 and state["processed_rows"] == 0
    monkeypatch.setattr(jobs, "iter_records", original)
    process_import(job)
    state = admin.get(f"/api/v1/imports/{job}").json()
    assert state["parsed_rows"] == state["processed_rows"] == state["total_rows"] == 260
    with get_sessionmaker()() as db:
        assert len(list(db.scalars(select(ImportRecord.id).where(ImportRecord.import_job_id == UUID(job))))) == 260


def test_dispatch_recovers_enqueue_outage_and_bounds_prestart_failures(actors, monkeypatch):
    from app.recruits import jobs
    from redis.exceptions import ConnectionError
    from rq.job import Job, JobStatus
    admin, _ = actors
    job = upload(admin, [row()])
    real_queue = jobs.queue
    def unavailable():
        raise ConnectionError("simulated queue outage")
    monkeypatch.setattr(jobs, "queue", unavailable)
    with pytest.raises(ConnectionError):
        dispatch_pending()
    assert admin.get(f"/api/v1/imports/{job}").json()["status"] == "QUEUED"
    monkeypatch.setattr(jobs, "queue", real_queue)
    for attempt in range(1, get_settings().job_max_attempts + 1):
        dispatch_pending()
        with get_sessionmaker()() as db:
            item = db.get(ImportJob, UUID(job))
            assert item.rq_job_id.endswith(f"-{attempt}")
            rq_job = Job.fetch(item.rq_job_id, connection=get_redis())
            queue().remove(rq_job.id)
            rq_job.set_status(JobStatus.FAILED)
    dispatch_pending()
    state = admin.get(f"/api/v1/imports/{job}").json()
    assert state["status"] == "FAILED"
    assert state["retry_count"] == get_settings().job_max_attempts


def test_recruit_documents_link_and_download(actors):
    admin, operator = actors
    recruit = approved_recruit(admin)
    other = approved_recruit(admin)
    response = admin.post(f"/api/v1/recruits/{recruit}/documents", files={"file": ("decision.pdf", b"%PDF-synthetic", "application/pdf")})
    assert response.status_code == 201, response.text
    document = response.json()["id"]
    assert operator.get(f"/api/v1/recruits/{recruit}/documents/{document}/download").content == b"%PDF-synthetic"
    assert operator.get(f"/api/v1/recruits/{other}/documents/{document}/download").status_code == 404


def test_label_preview_matches_print_document_without_creating_a_print_audit(actors):
    admin, operator = actors
    rid = approved_recruit(admin)
    assert operator.post(f"/api/v1/recruits/{rid}/complete", json=completion()).status_code == 200
    preview = operator.get(f"/api/v1/recruits/{rid}/barcode?preview=true")
    assert preview.status_code == 200
    with get_sessionmaker()() as db:
        assert not db.scalar(select(AuditLog.id).where(AuditLog.entity_id == rid, AuditLog.action == "label"))
    printed = operator.get(f"/api/v1/recruits/{rid}/barcode")
    assert printed.status_code == 200 and printed.content == preview.content
    with get_sessionmaker()() as db:
        entry = db.scalar(select(AuditLog).where(AuditLog.entity_id == rid, AuditLog.action == "label"))
        assert json.loads(entry.new_values)["event"] == "print_requested"


@pytest.mark.skipif(__import__('os').name != 'nt', reason="Microsoft ACE export requires a Windows worker")
def test_accdb_export_preserves_arabic(actors, monkeypatch):
    from app.recruits.files import path_for, access_connection, read_records
    admin, _ = actors
    job = upload(admin, [row(nid="bad")])
    process_import(job)
    response = admin.post(f"/api/v1/imports/{job}/exports", json={"format": "accdb"})
    export = response.json()["id"]
    process_export(export)
    with get_sessionmaker()() as db:
        item = db.get(ExportJob, UUID(export))
        assert item.status == "COMPLETED"
        with access_connection(path_for(item.storage_key)) as conn:
            result = conn.execute("SELECT [name], [rejection_reasons] FROM [rejected_records]").fetchone()
            assert result[0] == "أحمد محمد حسن علي"
            assert "الرقم القومي" in result[1]
            assert "(national_id):" in result[1]
        monkeypatch.setattr(get_settings(), "access_table", "rejected_records")
        imported = read_records(path_for(item.storage_key))
        assert len(imported) == 1 and imported[0]["name"] == "أحمد محمد حسن علي"
        assert imported[0]["national_id"] == "bad"


@pytest.mark.skipif(__import__('os').name != 'nt', reason="Native Windows worker regression")
def test_supervised_windows_worker_recovers_and_processes_import_and_export(actors, tmp_path):
    import os
    import subprocess
    import sys
    import time
    from app.core.config import REPO_ROOT
    admin, _ = actors
    job = upload(admin, [row(nid="bad")])
    settings = get_settings()
    env = dict(os.environ, DB_NAME=settings.db_name, QUEUE_NAME='poc-spawn-' + uuid4().hex,
               REDIS_PREFIX=settings.redis_prefix, STORAGE_ROOT=str(settings.storage_root),
               RECOVERY_INTERVAL_SECONDS='1')
    log_path = tmp_path / 'worker.log'
    with log_path.open('w', encoding='utf-8') as log:
        worker = subprocess.Popen([sys.executable, '-m', 'app.cli.worker_supervisor'], cwd=REPO_ROOT / 'backend',
            env=env, stdout=log, stderr=subprocess.STDOUT, creationflags=subprocess.CREATE_NO_WINDOW)
        try:
            def wait_for(path, expected):
                deadline = time.monotonic() + 30
                while time.monotonic() < deadline:
                    state = admin.get(path).json()
                    if state['status'] in {expected, 'FAILED'}:
                        break
                    time.sleep(.25)
                assert state['status'] == expected, log_path.read_text(encoding='utf-8')
                return state
            state = wait_for(f'/api/v1/imports/{job}', 'READY_FOR_REVIEW')
            assert state['invalid_rows'] == 1 and state['retry_count'] == 1
            # Crash the actual RQ worker, leaving its supervisor alive. The next
            # export must be completed by a new worker without manual intervention.
            import signal
            from rq import Worker
            from app.core.redis_client import get_redis
            running = [w for w in Worker.all(connection=get_redis()) if env['QUEUE_NAME'] in w.queue_names()]
            assert len(running) == 1
            os.kill(running[0].pid, signal.SIGTERM)
            export = admin.post(f'/api/v1/imports/{job}/exports', json={'format': 'xlsx'}).json()['id']
            wait_for(f'/api/v1/exports/{export}', 'COMPLETED')
            assert 'generation=2' in log_path.read_text(encoding='utf-8')
            assert admin.get(f'/api/v1/exports/{export}/download').content.startswith(b'PK')
            staged = admin.get(f'/api/v1/imports/{job}/report').json()['invalid']['items'][0]
            result = admin.patch(f'/api/v1/imports/{job}/rows/{staged["id"]}', json={
                'version': state['progress_version'], 'changes': {'national_id': row()['national_id']}})
            assert result.status_code == 202, result.text
            state = wait_for(f'/api/v1/imports/{job}', 'READY_FOR_REVIEW')
            assert state['valid_rows'] == 1 and state['invalid_rows'] == 0 and state['validation_run'] == 1
            assert admin.post(f'/api/v1/imports/{job}/revalidate', json={'version': state['progress_version']}).status_code == 202
            state = wait_for(f'/api/v1/imports/{job}', 'READY_FOR_REVIEW')
            assert state['validation_run'] == 2 and state['valid_rows'] == 1
        finally:
            subprocess.run(['taskkill', '/PID', str(worker.pid), '/T', '/F'], capture_output=True)
